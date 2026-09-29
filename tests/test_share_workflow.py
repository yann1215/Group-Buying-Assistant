import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from app.analysis.order_parser import parse_order_file
from app.analysis.product_config import (
    _read_product_config_rows, _write_product_config_rows,
    ensure_product_config_file, load_product_share_config_file,
    product_config_owner_path,
)
from app.core import path_manager as paths
from app.core.archive_manager import archive_conversation_files
from app.core.intent_parser import parse_user_intent
from app.core.tool_orchestrator import ToolOrchestrator, format_share_summary


class ShareIntentTests(unittest.TestCase):
    def test_input_is_not_calculation(self):
        for text in ('独立个数摊', '拉通人头120', '金额120', '1号80元，2号40元', '改成人头', '改成拉通'):
            with self.subTest(text=text):
                self.assertEqual(parse_user_intent(text)['intent'], 'update_share_config')

    def test_queries_are_read_only(self):
        for text in ('查看均摊', '看看均摊', '均摊是多少', '均摊多少', '查均摊', '看一下均摊', '查看均摊，金额120'):
            with self.subTest(text=text):
                intent = parse_user_intent(text)
                self.assertEqual(intent['intent'], 'show_share')
                self.assertIsNone(intent['amount'])

    def test_explicit_calculation(self):
        for text in ('算均摊', '计算均摊', '算一下均摊', '算均摊，拉通个数120', '算均摊，独立个数摊，1号80元，2号40元'):
            with self.subTest(text=text):
                self.assertEqual(parse_user_intent(text)['intent'], 'calculate_share')

    def test_only_one_dimension_is_changed(self):
        self.assertIsNone(parse_user_intent('改成人头摊')['calculation_scope'])
        self.assertIsNone(parse_user_intent('改成独立')['share_mode'])


class ShareWorkflowTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        for attr, directory in {
            'ORDER_INPUT_DIR': 'input', 'ORDER_CONFIG_DIR': 'config',
            'ORDER_OUTPUT_DIR': 'output', 'ORDER_ARCHIVE_DIR': 'archive',
            'WORKSPACE_DIR': 'workspace', 'TEMP_DIR': 'temp',
        }.items():
            self.enterContext(patch.object(paths, attr, self.root / directory))
        self.order = self.root / 'orders.xlsx'
        wb = Workbook()
        ws = wb.create_sheet('订单')
        ws.append([])
        ws.append([])
        ws.append(['单号', '昵称', '总金额', '发货状态', '商品A', '商品B', '底胚'])
        for row in [(1, '甲', 10, '', 2, 1, 0), (2, '乙', 10, '', 1, 0, 1),
                    (3, '丙', 10, '', 0, 3, 0), (4, '车主', 10, '', 10, 10, 0)]:
            ws.append([value if value != 0 else None for value in row])
        wb.save(self.order)
        wb.close()
        self.tools = ToolOrchestrator()
        self.ctx = self.tools.get_context(1)
        self.ctx.group_name = '测试群'
        self.ctx.new_order_file = str(self.order)
        self.ctx.special_members = [{'角色': '车主', '昵称': '车主', '单号': '4', '参摊': False}]
        self.members = self.enterContext(patch(
            'app.core.tool_orchestrator.parse_group_member_orders', side_effect=self.member_result))

    def member_result(self, **kwargs):
        parsed = parse_order_file(kwargs['order_input'], kwargs['parsed_output_path'])
        config = ensure_product_config_file(parsed, kwargs['group_name'])
        return {'ok': True, '群聊名称': kwargs['group_name'], 'member_count': 4,
                'parsed_order_file': parsed, 'share_config_file': config,
                'product_configs': load_product_share_config_file(config)}

    def say(self, text):
        return self.tools.handle(1, text)

    def configs(self):
        return load_product_share_config_file(self.ctx.share_config_file)

    def edit_config(self, index, **changes):
        rows = _read_product_config_rows(self.ctx.share_config_file)
        rows[index].update(changes)
        _write_product_config_rows(Path(self.ctx.share_config_file), rows)

    def calculate(self, mode='拉通个数摊120'):
        self.say(mode)
        if '独立' in mode:
            self.say('1号80元，2号40元')
        self.assertIn('请确认商品配置', self.say('算均摊'))
        reply = self.say('确认计算')
        self.assertIn('实际总收款', reply)
        return reply

    def test_input_and_missing_amount_never_check_members(self):
        self.assertIn('未执行计算', self.say('独立个数摊'))
        self.assertIn('请补充', self.say('算均摊'))
        self.assertIn('没有待确认', self.say('确认计算'))
        self.members.assert_not_called()
        self.assertFalse(paths.ORDER_OUTPUT_DIR.exists())

    def test_confirmation_precedes_members_in_all_modes(self):
        for mode in ('拉通人头摊120', '拉通个数摊120', '独立人头摊', '独立个数摊'):
            with self.subTest(mode=mode):
                self.members.reset_mock()
                self.say(mode)
                if '独立' in mode:
                    self.say('1号80元，2号40元')
                preview = self.say('算均摊')
                self.assertIn('底胚：不参与均摊', preview)
                self.members.assert_not_called()
                reply = self.say('确认计算')
                self.assertIn('实际总收款', reply)
                self.members.assert_called_once()
                result = self.ctx.last_share_result
                self.assertEqual(result['participant_count'], 3)
                self.assertEqual([c['quantity'] for c in result['product_statistics']], [3, 4, 0])
                self.assertIn(format_share_summary(result), reply)
                self.assertEqual(self.say('查看均摊'), format_share_summary(result))

    def test_independent_head_statistics(self):
        reply = self.calculate('独立人头摊')
        self.assertIn('参摊人数：3 人', reply)
        self.assertIn('商品A：2 人', reply)
        self.assertIn('商品B：2 人', reply)
        self.assertIn('商品A：40.00', reply)
        self.assertIn('商品B：20.00', reply)

    def test_query_has_no_writes_or_member_calls(self):
        self.calculate()
        before = {p: p.read_bytes() for p in self.root.rglob('*') if p.is_file()}
        self.members.reset_mock()
        for text in ('查看均摊', '看看均摊', '均摊是多少'):
            self.assertIn('参摊个数：7 个', self.say(text))
        self.members.assert_not_called()
        self.assertEqual(before, {p: p.read_bytes() for p in self.root.rglob('*') if p.is_file()})

    def test_show_before_calculation(self):
        self.say('独立个数摊')
        reply = self.say('查看均摊')
        self.assertIn('待计算', reply)
        self.assertIn('未填写', reply)
        self.members.assert_not_called()

    def test_switch_clears_old_amounts_and_preserves_other_fields(self):
        self.calculate()
        self.edit_config(0, 商品单价='10.00', 商品大货总价='130.00')
        self.say('改成独立')
        rows = self.configs()
        self.assertTrue(all(c['商品均摊'] in ('', '0.00') and c['单份均摊'] == '' for c in rows))
        self.assertEqual(rows[0]['商品单价'], '10.00')
        self.assertEqual(rows[0]['商品大货总价'], '130.00')
        self.assertFalse(rows[2]['计入均摊'])
        self.assertEqual(self.ctx.share_request.amount, '120')
        self.assertIn('失效', self.say('查看均摊'))
        self.assertIn('请补充', self.say('算均摊'))

    def test_switch_independent_mode_clears_and_new_input_wins(self):
        self.calculate('独立个数摊')
        self.say('人头摊，1号60元，2号40元')
        self.assertEqual(self.ctx.share_request.calculation_scope, 'independent')
        self.assertEqual([c['商品均摊'] for c in self.configs()[:2]], ['60.00', '40.00'])
        self.assertTrue(all(c['单份均摊'] == '' for c in self.configs()))

    def test_same_mode_preserves_amounts_and_unit_prices(self):
        self.calculate('独立个数摊')
        before = self.configs()
        self.say('独立个数摊')
        self.assertEqual(before, self.configs())
        self.assertNotIn('失效', self.say('查看均摊'))

    def test_switch_to_flat_rewrites_total(self):
        self.calculate('拉通个数摊120')
        self.say('独立个数摊，1号80元，2号40元')
        self.say('改成拉通')
        self.assertTrue(all(c['商品均摊'] == '120.00' for c in self.configs()))
        self.assertTrue(all(c['单份均摊'] == '' for c in self.configs()))

    def test_switch_preserves_implicit_independent_total(self):
        self.calculate('独立个数摊')
        self.say('改成拉通')
        self.assertEqual(self.ctx.share_request.amount, '120.00')
        self.assertTrue(all(c['商品均摊'] == '120.00' for c in self.configs()))

    def test_amount_mismatch_blocks_before_members(self):
        self.say('独立个数摊，总均摊120，1号80元，2号30元')
        self.assertIn('不一致', self.say('算均摊'))
        self.members.assert_not_called()

    def test_combined_command_still_requires_confirmation(self):
        self.assertIn('请确认商品配置', self.say('算均摊，独立个数摊，1号80元，2号40元'))
        self.members.assert_not_called()
        self.assertIn('实际总收款', self.say('确认计算'))

    def test_force_cannot_bypass_confirmation(self):
        self.say('拉通个数摊120')
        self.assertIn('请确认商品配置', self.say('忽略名单问题，继续计算'))
        self.members.assert_not_called()

    def test_edit_config_after_preview_requires_new_confirmation(self):
        self.say('拉通个数摊120')
        self.say('算均摊')
        self.edit_config(1, 计入均摊=False)
        self.assertIn('重新确认', self.say('确认计算'))
        self.members.assert_not_called()
        self.assertIn('实际总收款', self.say('确认计算'))

    def test_order_overwritten_after_preview_requires_new_confirmation(self):
        self.say('拉通个数摊120')
        self.say('算均摊')
        wb = load_workbook(self.order)
        wb.worksheets[1].cell(4, 5, 3)
        wb.save(self.order)
        wb.close()
        self.assertIn('重新确认', self.say('确认计算'))
        self.members.assert_not_called()

    def test_member_sync_changes_config_and_stops_calculation(self):
        self.say('拉通个数摊120')
        self.say('算均摊')
        def changed(**kwargs):
            result = self.member_result(**kwargs)
            self.edit_config(1, 计入均摊=False)
            return result
        self.members.side_effect = changed
        self.assertIn('暂不计算', self.say('确认计算'))
        self.assertIsNone(self.ctx.last_share_result)

    def test_changes_invalidate_result_and_restore_is_supported(self):
        self.calculate()
        saved = json.loads(json.dumps(self.ctx.to_dict()))
        self.tools.load_context(1, saved)
        self.ctx = self.tools.get_context(1)
        self.assertNotIn('失效', self.say('看看均摊'))
        self.ctx.special_members[0]['参摊'] = True
        self.assertIn('失效', self.say('查看均摊'))

    def test_same_group_new_session_does_not_reuse_amounts(self):
        self.calculate('独立个数摊')
        old_token = self.ctx.config_owner_id
        self.tools.remove_context(1)
        self.ctx = self.tools.get_context(1)  # 即使数据库编号复用，归属 token 也不同。
        self.ctx.group_name = '测试群'
        self.ctx.new_order_file = str(self.order)
        self.say('独立个数摊')
        self.assertNotEqual(old_token, self.ctx.config_owner_id)
        self.assertTrue(all(c['商品均摊'] in ('', '0.00') for c in self.configs()))
        self.assertIsNone(self.ctx.share_request.amount)
        self.assertTrue(list((paths.ORDER_ARCHIVE_DIR / 'unclaimed').rglob('*.csv')))

    def test_delete_archives_config_marker_and_new_session_is_clean(self):
        self.calculate()
        config = Path(self.ctx.share_config_file)
        with archive_conversation_files(1, self.ctx.group_name, self.ctx.to_dict()) as archive:
            pass
        self.assertFalse(config.exists())
        self.assertFalse(product_config_owner_path(config).exists())
        self.assertTrue(list((archive / 'config').glob('*.owner.json')))
        self.tools.remove_context(1)
        self.ctx = self.tools.get_context(2)
        self.ctx.group_name = '测试群'
        self.ctx.new_order_file = str(self.order)
        self.tools.handle(2, '独立个数摊')
        self.assertIsNone(self.ctx.share_request.amount)
        self.assertTrue(all(c['商品均摊'] in ('', '0.00') for c in self.configs()))

    def test_zero_independent_amount_still_counts_participants(self):
        self.say('独立人头摊，1号0元，2号0元')
        self.say('算均摊')
        reply = self.say('确认计算')
        self.assertIn('参摊人数：3 人', reply)
        self.assertEqual(self.ctx.last_share_result['total_collected'], '0.00')

    def test_confirmation_with_new_parameters_requires_another_confirmation(self):
        self.say('拉通个数摊120')
        self.say('算均摊')
        self.assertIn('请确认商品配置', self.say('确认计算，金额140'))
        self.members.assert_not_called()
        self.assertIn('总均摊：140.00', self.say('确认计算'))

    def test_negative_confirmation_never_calculates(self):
        for text in ('暂不确认', '不确认计算', '取消', '不要计算均摊',
                     '不要确认配置', '配置有问题，不确认', '先不计算'):
            self.say('拉通个数摊120')
            self.say('算均摊')
            self.assertIn('已取消', self.say(text))
        self.members.assert_not_called()

    def test_force_continues_only_the_confirmed_request(self):
        def blocked(**kwargs):
            result = self.member_result(**kwargs)
            result['serials_in_orders_not_in_group'] = ['1']
            return result
        self.members.side_effect = blocked
        self.say('拉通个数摊120')
        self.say('算均摊')
        self.assertIn('名单核对问题', self.say('确认计算'))
        self.assertIsNone(self.ctx.last_share_result)
        self.assertIn('实际总收款', self.say('忽略名单问题，继续计算'))
        self.assertFalse(self.ctx.share_request.force)
        self.assertIn('请确认商品配置', self.say('忽略名单问题，继续计算'))

    def test_external_config_and_order_changes_invalidate_query(self):
        self.calculate()
        self.edit_config(0, 商品均摊='130.00')
        self.assertIn('失效', self.say('查看均摊'))
        self.calculate()
        wb = load_workbook(self.order)
        wb.worksheets[1].cell(4, 5, 3)
        wb.save(self.order)
        wb.close()
        self.assertIn('失效', self.say('查看均摊'))

    def test_legacy_context_adopts_its_config_without_losing_amounts(self):
        self.calculate('独立个数摊')
        before = self.configs()
        saved = self.ctx.to_dict()
        saved.pop('config_owner_id')
        product_config_owner_path(self.ctx.share_config_file).unlink()
        self.tools.load_context(1, saved)
        self.ctx = self.tools.get_context(1)
        self.assertEqual(before, self.configs())
        self.assertTrue(product_config_owner_path(self.ctx.share_config_file).exists())

    def test_unowned_legacy_file_is_backed_up_for_new_conversation(self):
        self.calculate('独立个数摊')
        product_config_owner_path(self.ctx.share_config_file).unlink()
        self.tools.remove_context(1)
        self.ctx = self.tools.get_context(1)
        self.ctx.group_name = '测试群'
        self.ctx.new_order_file = str(self.order)
        self.assertIn('请补充', self.say('独立个数摊'))
        self.assertTrue(list((paths.ORDER_ARCHIVE_DIR / 'unclaimed').rglob('*.csv')))

    def test_rename_preserves_ownership_and_delete_rolls_back(self):
        self.calculate()
        self.tools.set_context(1, group_name='改名后的群')
        marker = product_config_owner_path(self.ctx.share_config_file)
        self.assertEqual(json.loads(marker.read_text())['owner_id'], self.ctx.config_owner_id)
        with self.assertRaises(RuntimeError):
            with archive_conversation_files(1, self.ctx.group_name, self.ctx.to_dict()):
                raise RuntimeError('模拟数据库删除失败')
        self.assertTrue(Path(self.ctx.share_config_file).exists())
        self.assertTrue(marker.exists())
        self.assertTrue(Path(self.ctx.parsed_order_file).exists())

    def test_chat_service_delete_and_recreate_same_group(self):
        from app.database import db, repositories
        from app.core.chat_service import ChatService

        self.enterContext(patch.object(db, 'DB_PATH', self.root / 'app.db'))
        self.enterContext(patch.object(db, 'ensure_dirs'))
        connect = db.sqlite3.connect
        def tracked_connect(*args, **kwargs):
            connection = connect(*args, **kwargs)
            self.addCleanup(connection.close)
            return connection
        self.enterContext(patch.object(db.sqlite3, 'connect', side_effect=tracked_connect))
        db.init_db()
        service = ChatService()
        session = service.create_conversation(group_name='生命周期测试群')
        service.set_working_context(session, order_input=self.order)
        service.send_message(session, '独立个数摊，1号80元，2号40元')
        old = service.tools.get_context(session)
        old_config = Path(old.share_config_file)
        self.assertEqual(load_product_share_config_file(old_config)[0]['商品均摊'], '80.00')
        service.delete_conversation(session)
        self.assertFalse(old_config.exists())
        self.assertFalse(product_config_owner_path(old_config).exists())
        self.assertIsNone(repositories.get_session(session))
        new_session = service.create_conversation(group_name='生命周期测试群')
        service.set_working_context(new_session, order_input=self.order)
        reply = service.send_message(new_session, '独立个数摊')
        self.assertIn('请补充', reply)
        self.members.assert_not_called()

    def test_relaxed_confirmation_words(self):
        for text in ('计算', '算', '算吧', '开始计算', '无误', '下一步', '继续', '没问题', '可以', '算均摊'):
            with self.subTest(text=text):
                self.say('拉通人头摊50')
                self.say('算均摊')
                self.members.reset_mock()
                reply = self.say(text)
                self.assertIn('实际总收款', reply)
                self.members.assert_called_once()

    def test_relaxed_words_require_pending_confirmation(self):
        self.say('拉通人头摊50')
        for text in ('算', '计算', '无误', '下一步', '可以'):
            self.assertIsNone(self.say(text))
        self.members.assert_not_called()

    def test_relaxed_confirmation_does_not_ignore_members(self):
        def blocked(**kwargs):
            result = self.member_result(**kwargs)
            result['serials_in_orders_not_in_group'] = ['1']
            return result
        self.members.side_effect = blocked
        self.say('拉通人头摊50')
        self.say('算均摊')
        self.assertIn('名单核对问题', self.say('继续算'))
        self.assertIsNone(self.ctx.last_share_result)

    def test_relaxed_confirmation_with_changes_repreviews(self):
        self.say('拉通人头摊50')
        self.say('算均摊')
        self.assertIn('请确认商品配置', self.say('下一步，金额60'))
        self.members.assert_not_called()
        self.assertIn('总均摊：60.00', self.say('算'))

    def test_summary_format_for_query_and_preview(self):
        self.say('拉通人头摊50')
        preview = self.say('算均摊')
        self.assertIn('均摊类型：人头摊\n计算方式：拉通\n总均摊：50.00', preview)
        reply = self.say('算')
        summary = self.say('查看均摊')
        self.assertIn(summary, reply)
        self.assertIn('参摊人数：3 人\n单人均摊：16.67', summary)
        self.assertNotIn('去重', summary)
        self.assertNotIn('元', summary)

    def restore_legacy(self, mode='拉通人头摊50'):
        self.calculate(mode)
        saved = self.ctx.to_dict()
        for key in ('last_share_result', 'last_share_signature', 'legacy_share_signature', 'share_results_invalidated'):
            saved.pop(key, None)
        saved['share_request'] = {}  # 方式和总额也必须能从 CSV 读取。
        self.ctx = self.tools.load_context(1, saved)
        self.members.reset_mock()

    def test_legacy_flat_csv_is_displayed_without_recalculation(self):
        self.restore_legacy()
        before = {p: p.read_bytes() for p in self.root.rglob('*') if p.is_file()}
        reply = self.say('查看均摊')
        self.assertIn('历史配置记录', reply)
        self.assertIn('均摊类型：人头摊\n计算方式：拉通\n总均摊：50.00', reply)
        self.assertIn('参摊人数：历史记录未保存', reply)
        self.assertIn('单人均摊：16.67', reply)
        self.assertNotIn('尚未计算', reply)
        self.members.assert_not_called()
        self.assertEqual(before, {p: p.read_bytes() for p in self.root.rglob('*') if p.is_file()})

    def test_legacy_independent_csv_uses_saved_amounts_not_product_quantity(self):
        self.restore_legacy('独立个数摊')
        reply = self.say('看看均摊')
        self.assertIn('总均摊：120.00', reply)
        self.assertIn('商品A独立均摊：80.00', reply)
        self.assertIn('商品A：26.67', reply)
        self.assertIn('参摊个数：历史记录未保存', reply)
        self.assertNotIn('参摊个数：27', reply)

    def test_legacy_edit_invalidates_fallback_even_after_reload(self):
        self.restore_legacy()
        saved = self.ctx.to_dict()
        self.edit_config(0, 计入均摊=False)
        self.ctx = self.tools.load_context(1, saved)
        reply = self.say('查看均摊')
        self.assertIn('失效', reply)
        self.assertNotIn('单人均摊：16.67', reply)

    def test_explicitly_invalidated_legacy_does_not_revive_csv(self):
        self.restore_legacy()
        from app.core.tool_orchestrator import invalidate_share_confirmation
        invalidate_share_confirmation(self.ctx)
        self.ctx = self.tools.load_context(1, self.ctx.to_dict())
        self.assertIn('失效', self.say('查看均摊'))

    def test_legacy_without_unit_prices_remains_pending(self):
        self.say('独立个数摊，1号80元，2号40元')
        saved = self.ctx.to_dict()
        saved.pop('share_results_invalidated')
        saved['share_request'] = {}
        self.ctx = self.tools.load_context(1, saved)
        reply = self.say('查看均摊')
        self.assertIn('尚未计算', reply)
        self.assertIn('待计算', reply)
        self.assertIn('均摊类型：个数摊\n计算方式：独立\n总均摊：120.00', reply)

    def test_order_path_display_in_all_slots_and_errors(self):
        from app.core.chat_service import ChatService
        from app.core.order_version_manager import OrderVersionUpdateResult, RemovedOrderPath
        from app.core.tool_orchestrator import format_context_update_result
        inside = paths.ORDER_INPUT_DIR / '子目录' / '订单.xlsx'
        outside = self.root / 'input_backup' / '订单.xlsx'
        self.assertEqual(paths.format_order_path(inside), '订单.xlsx')
        self.assertEqual(paths.format_order_path(outside), str(outside.resolve()))
        self.assertEqual(paths.format_order_path('订单.xlsx'), '订单.xlsx')
        self.assertEqual(paths.format_order_path(None), '未设置')
        versions = {'new_order_file': str(inside), 'old_order_file': str(outside),
                    'order_cache_1_file': str(inside), 'order_cache_2_file': ''}
        result = OrderVersionUpdateResult(versions=versions, success=False, input_path=str(inside),
                                         removed_paths=(RemovedOrderPath('new_order_file', str(inside), '无效'),))
        reply = ChatService._format_order_update_result(result)
        self.assertIn('已检查：订单.xlsx', reply)
        self.assertIn('新订单：订单.xlsx', reply)
        self.assertNotIn(str(inside), reply)
        self.assertIn(str(outside.resolve()), reply)
        for key, value in versions.items():
            setattr(self.ctx, key, value)
        self.assertIn('新订单：订单.xlsx', format_context_update_result(self.ctx))


if __name__ == '__main__':
    unittest.main()
