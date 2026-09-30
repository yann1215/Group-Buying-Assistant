import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

import test_share_workflow as fixtures
from app.analysis.participation import parse_participation, confirmation
from app.analysis.product_config import update_product_share_config_file, ensure_product_config_file
from app.analysis.order_validator import inspect_order_status
from app.core.tool_orchestrator import SessionToolContext
from app.utils.name_matching import match_names


class ParticipationTests(unittest.TestCase):
    setUp = fixtures.ShareWorkflowTests.setUp
    member_result = fixtures.ShareWorkflowTests.member_result
    say = fixtures.ShareWorkflowTests.say
    configs = fixtures.ShareWorkflowTests.configs

    def test_product_saves_immediately_and_previews_only_target(self):
        self.say('拉通个数摊120')
        reply = self.say('商品底胚参摊')
        self.assertIn('已保存', reply)
        self.assertIn('底胚', reply)
        self.assertNotIn('商品A', reply)
        self.assertNotIn('商品B', reply)
        self.assertTrue(self.configs()[2]['计入均摊'])
        self.assertIsNone(self.ctx.pending_participation)
        self.say('先别改')
        self.assertTrue(self.configs()[2]['计入均摊'])
        self.members.assert_not_called()

    def test_product_first_and_explicit_member(self):
        self.say('拉通个数摊120')
        self.ctx.special_members.append({'角色': '工具人', '昵称': '底胚', '单号': '2', '参摊': True})
        self.say('底胚参摊')
        self.assertTrue(self.configs()[2]['计入均摊'])
        self.assertTrue(self.ctx.special_members[-1]['参摊'])
        reply = self.say('成员底胚不参摊')
        self.assertIn('成员：底胚', reply)
        self.assertNotIn('车主', reply)
        self.assertFalse(self.ctx.special_members[-1]['参摊'])

    def test_ordinary_member_and_existing_role(self):
        original = deepcopy(self.ctx.special_members)
        self.assertIn('其他不参摊成员', self.say('成员甲不参摊'))
        self.assertEqual(self.ctx.special_members[0], original[0])
        self.assertEqual(self.ctx.special_members[-1]['单号'], '1')
        self.assertFalse(self.ctx.special_members[-1]['参摊'])
        self.say('成员车主参摊')
        self.say('OK')
        self.assertEqual(self.ctx.special_members[0]['角色'], '车主')
        self.assertTrue(self.ctx.special_members[0]['参摊'])

    def test_bare_member_fallback(self):
        self.assertIn('成员：甲', self.say('甲不参摊'))

    def test_explicit_product_does_not_fall_back(self):
        self.assertIn('未找到匹配商品', self.say('商品甲不参摊'))
        self.assertIsNone(self.ctx.pending_participation)

    def test_product_name_starting_with_product_prefix(self):
        self.assertIn('商品A', self.say('商品A不参摊'))
        self.say('好的')
        self.assertFalse(self.configs()[0]['计入均摊'])

    def test_duplicate_group_members_require_selection(self):
        group = {'ok': True, 'members': [
            {'昵称': '小王甲', '群昵称': '11 小王甲', 'wxid': 'a'},
            {'昵称': '小王乙', '群昵称': '12 小王乙', 'wxid': 'b'},
        ]}
        with patch('app.core.participation_workflow.get_wechat_group_members', return_value=group):
            self.assertIn('多个候选', self.say('成员小王不参摊'))
        self.assertEqual(len(self.ctx.special_members), 1)
        reply = self.say('选择2')
        self.assertIn('已保存', reply)
        self.assertNotIn('小王甲', reply)
        self.assertEqual(self.ctx.special_members[-1]['wxid'], 'b')
        self.assertEqual(self.ctx.special_members[-1]['单号'], '12')

    def test_cancel_selection_preserves_valid_results(self):
        self.say('拉通个数摊120')
        self.say('算均摊')
        self.say('确认计算')
        signature = self.ctx.last_share_signature
        self.say('商品商品不参摊')
        self.say('取消')
        self.assertEqual(self.ctx.last_share_signature, signature)
        self.assertFalse(self.ctx.share_results_invalidated)

    def test_member_change_changes_calculation(self):
        self.say('拉通个数摊120')
        self.say('成员甲不参摊')
        self.say('确认')
        self.say('算均摊')
        self.say('确认计算')
        self.assertEqual(self.ctx.last_share_result['total_share_quantity'], 4)
        self.assertNotIn('1', [row['单号'] for row in self.ctx.last_share_result['items']])

    def test_multiple_products_never_fall_back(self):
        self.assertIn('多个候选', self.say('商品商品不参摊'))
        self.assertIn('多个候选', self.say('好的'))
        reply = self.say('选择2')
        self.assertIn('商品B', reply)
        self.assertIn('已保存', reply)
        self.assertNotIn('商品A', reply)
        self.assertFalse(self.configs()[1]['计入均摊'])
        self.assertTrue(self.configs()[0]['计入均摊'])

    def test_correction_selects_unique_target_and_saves(self):
        self.say('商品商品不参摊')
        self.assertIn('商品B', self.say('可以，但改成商品B'))
        self.assertFalse(self.configs()[1]['计入均摊'])
        self.assertTrue(self.configs()[0]['计入均摊'])
        self.assertIsNone(self.ctx.pending_participation)

    def test_correction_with_state(self):
        self.say('商品商品不参摊')
        self.assertIn('底胚', self.say('可以，但改成商品底胚参摊'))
        self.assertTrue(self.configs()[2]['计入均摊'])
        self.assertIsNone(self.ctx.pending_participation)

    def test_selection_survives_serialization(self):
        self.say('商品商品不参摊')
        self.tools.contexts[1] = SessionToolContext.from_dict(self.ctx.to_dict())
        self.tools.contexts[1].session_id = 1
        self.ctx = self.tools.contexts[1]
        self.assertIn('已保存', self.say('选择2'))
        self.assertFalse(self.configs()[1]['计入均摊'])

    def test_external_edit_invalidates_selection(self):
        self.say('商品商品不参摊')
        path = Path(self.ctx.share_config_file)
        path.write_bytes(path.read_bytes() + b'\n')
        self.assertIn('已变化', self.say('选择2'))
        self.assertTrue(self.configs()[1]['计入均摊'])

    def test_confirmation_does_not_calculate(self):
        self.say('拉通个数摊120')
        self.say('算均摊')
        self.say('底胚参摊')
        self.say('继续')
        self.assertFalse(self.ctx.share_request.pending_config_confirmation)
        self.members.assert_not_called()

    def test_legacy_member_command_saves_immediately(self):
        self.assertIn('已保存', self.say('车主：昵称=小王，单号=4，参摊'))
        self.assertTrue(self.ctx.special_members[0]['参摊'])
        self.assertIsNone(self.ctx.pending_participation)

    def test_old_confirmation_draft_is_never_executed(self):
        for draft in ({'kind': 'legacy', 'members': []},
                      {'kind': '商品', 'candidates': [], 'include': False}):
            self.ctx.pending_participation = draft
            restored = SessionToolContext.from_dict(self.ctx.to_dict())
            restored.session_id = 1
            self.ctx = self.tools.contexts[1] = restored
            self.assertIn('旧版待确认修改已清除', self.say('确认'))
            self.assertEqual(len(self.ctx.special_members), 1)
            self.assertIsNone(self.ctx.pending_participation)

    def test_amount_exact_fuzzy_and_ambiguous(self):
        self.say('独立个数摊')
        result = update_product_share_config_file(self.ctx.share_config_file, [{'商品名称': '品A', '商品均摊': '80'}])
        self.assertEqual(len(result['updated_items']), 1)
        before = Path(self.ctx.share_config_file).read_bytes()
        result = update_product_share_config_file(self.ctx.share_config_file, [{'商品名称': '商品', '商品均摊': '20'}])
        self.assertEqual(len(result['unmatched_updates'][0]['候选商品']), 2)
        self.assertEqual(Path(self.ctx.share_config_file).read_bytes(), before)

    def test_amount_preview_is_local_but_calculation_preview_is_complete(self):
        self.say('独立个数摊')
        self.say('1号80元，2号40元')
        reply = self.say('1号60元')
        self.assertIn('商品A', reply)
        self.assertIn('60.00', reply)
        self.assertNotIn('商品B', reply)
        self.assertNotIn('底胚', reply)
        self.assertNotIn('车主', reply)
        self.say('画师 小李')
        for text in ('算均摊', '算大货'):
            preview = self.say(text)
            for name in ('商品A', '商品B', '底胚', '车主', '小李'):
                self.assertIn(name, preview)
            self.assertIn('请确认', preview)
        self.members.assert_not_called()

    def test_old_draft_cleanup_keeps_calculation_confirmation(self):
        self.say('拉通个数摊120')
        self.say('算均摊')
        self.ctx.pending_participation = {'kind': 'legacy', 'members': []}
        self.assertIn('旧版待确认修改已清除', self.say('确认'))
        self.assertTrue(self.ctx.share_request.pending_config_confirmation)
        self.members.assert_not_called()
        self.assertEqual(len(self.ctx.special_members), 1)

    def test_config_override_persists_and_calculates(self):
        self.say('拉通个数摊120')
        self.say('底胚参摊')
        self.say('好')
        ensure_product_config_file(self.ctx.parsed_order_file, self.ctx.group_name)
        self.assertTrue(self.configs()[2]['计入均摊'])
        self.say('算均摊')
        self.assertIn('实际总收款', self.say('确认计算'))

    def test_special_product_override(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.order)
        wb['订单'].cell(3, 7, '画师专拍')
        wb.save(self.order)
        wb.close()
        self.say('拉通个数摊120')
        self.say('商品画师专拍参摊')
        self.say('确认')
        ensure_product_config_file(self.ctx.parsed_order_file, self.ctx.group_name)
        configs = self.configs()
        self.assertTrue(configs[2]['计入均摊'])
        self.assertFalse(inspect_order_status(self.ctx.parsed_order_file, configs)['special_product_orders'])
        from app.analysis.share_calculator import calculate_share
        result = calculate_share(parsed_order_file=self.ctx.parsed_order_file,
                                 group_name=self.ctx.group_name, share_mode='quantity',
                                 calculation_scope='flat', total_amount=120,
                                 product_configs=configs, excluded_order_nos={'4'})
        self.assertTrue(result['ok'])
        self.assertEqual(result['total_share_quantity'], 8)
        for config in configs:
            config['均摊类型'] = '独立个数摊'
            config['商品均摊'] = '40'
        result = calculate_share(parsed_order_file=self.ctx.parsed_order_file,
                                 group_name=self.ctx.group_name, share_mode='quantity',
                                 calculation_scope='independent', total_amount=120,
                                 product_configs=configs, excluded_order_nos={'4'})
        special = next(row for row in result['product_statistics'] if row['product_name'] == '画师专拍')
        self.assertTrue(special['included'])
        self.assertEqual(special['quantity'], 1)


class MatchingTests(unittest.TestCase):
    def test_exact_precedes_fuzzy(self):
        rows = [{'name': '小猫'}, {'name': '小猫挂件'}]
        self.assertEqual(match_names('小猫', rows, ['name']), rows[:1])

    def test_commands_and_confirmation(self):
        self.assertEqual(parse_participation('商品小猫不参摊')['query'], '小猫')
        self.assertFalse(parse_participation('商品小猫不参摊')['include'])
        self.assertFalse(confirmation('不确认'))
        self.assertFalse(confirmation('可以，但是改成小王'))
        for value in ('确认', '确定', '对', '是的', '没错', '正确', '可以', '好的', '好', 'OK', '没问题', '就这样', '执行', '继续'):
            self.assertTrue(confirmation(value))
