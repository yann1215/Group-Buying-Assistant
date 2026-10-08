"""合发录入回归检查：使用临时数据库和订单，不访问真实微信数据。

运行：python -m scripts.verify_merge_registration
"""
from contextlib import ExitStack
from pathlib import Path
from tempfile import TemporaryDirectory
import gc
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from app.core.chat_service import ChatService
from app.core.intent_parser import parse_user_intent
from app.core.session_types import MERGED_SHIPPING, SINGLE_CAR, UNCLASSIFIED
from app.database import db
from app.database.repositories import get_session


class MergeRegistrationChecks(unittest.TestCase):
    def setUp(self):
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.root = Path(self.stack.enter_context(TemporaryDirectory()))
        self.stack.callback(gc.collect)
        for target, value in (
            ("app.database.db.DB_PATH", self.root / "test.db"),
            ("app.core.path_manager.ORDER_INPUT_DIR", self.root),
            ("app.core.order_merge_workflow.ORDER_OUTPUT_DIR", self.root),
            ("app.core.path_manager.WORKSPACE_DIR", self.root / "workspace"),
        ):
            self.stack.enter_context(patch(target, value))
        self.stack.enter_context(patch("app.database.db.ensure_dirs"))
        db.init_db()
        for name in ("order1", "order2"):
            workbook = Workbook()
            workbook.save(self.root / f"{name}.xlsx")
            workbook.close()
        self.service = ChatService()
        self.stack.enter_context(patch.object(
            self.service.instruction_normalizer, "normalize",
            side_effect=AssertionError("精确录入指令不应调用模型")))
        self.session_id = self.service.create_conversation()

    def send(self, text):
        return self.service.send_message(self.session_id, text)

    def register(self):
        self.send("合发")
        self.send("车1 aaa，订单 order1")
        self.send("车群 bbb，文件 order2")

    def test_requested_dialogue_and_reload(self):
        self.assertEqual(self.send("合发"),
                         "当前对话类型已设置为合发对话，请录入车群信息。\n示例：车名 xxx，订单 xxx")
        self.assertEqual(get_session(self.session_id)["session_type"], MERGED_SHIPPING)
        self.assertEqual(self.send("车1 aaa，订单 order1"),
                         "已录入信息：车群1 aaa，订单 order1\n当前合发车群数量：1，请继续录入车群信息")
        self.assertIn('如需计算合发补邮清单，请输入指令“输出合发表”', self.send("车群 bbb，文件 order2"))
        self.service.tools.remove_context(self.session_id)
        self.service.load_conversation(self.session_id)
        ctx = self.service.tools.get_context(self.session_id)
        self.assertEqual(ctx.merge_groups, ["aaa", "bbb"])
        self.assertEqual(ctx.merge_order_files["aaa"], str(self.root / "order1.xlsx"))
        self.assertEqual(self.send("查看合发订单"), "1. aaa：order1.xlsx\n2. bbb：order2.xlsx")
        self.send("合发")
        self.assertEqual(ctx.merge_groups, ["aaa", "bbb"])

    def test_invalid_and_duplicate_entries(self):
        self.send("合发")
        self.assertIn("请同时录入", self.send("车名 aaa"))
        self.assertIn("订单无效", self.send("车名 aaa，订单 missing"))
        ctx = self.service.tools.get_context(self.session_id)
        self.assertEqual(ctx.merge_groups, [])
        self.send("车名：aaa，文件：order1")
        self.assertIn("不能重复", self.send("车群 aaa，订单 order2"))
        self.assertEqual(ctx.merge_groups, ["aaa"])
        self.assertIn("至少录入两个", self.send("输出合发表"))

    def test_type_isolation_and_legacy_format(self):
        self.assertEqual(parse_user_intent("合发", UNCLASSIFIED)["intent"], "start_merged_shipping")
        self.assertEqual(parse_user_intent("合发", SINGLE_CAR)["intent"], "unsupported")
        self.assertEqual(parse_user_intent("合发：aaa，bbb", MERGED_SHIPPING)["merge_groups"], ["aaa", "bbb"])
        self.register()
        self.assertIn("不支持", self.send("算均摊"))
        self.send("合发：ccc，ddd")
        ctx = self.service.tools.get_context(self.session_id)
        self.assertEqual(ctx.merge_groups, ["ccc", "ddd"])
        self.assertEqual(ctx.merge_order_files, {})

    def test_direct_orders_export_and_cache(self):
        self.register()
        def orders(path, group):
            self.assertEqual(Path(path), self.root / ("order1.xlsx" if group == "aaa" else "order2.xlsx"))
            return [{"serial": "1", "name": group, "row": 4, "products": {"徽章": 1},
                     "shipping": [group, "123", "地址"]}]
        def members(**kwargs):
            group = kwargs["group_name"]
            return {"ok": True, "members": [{"wxid": f"wxid_{group}", "昵称": group, "群昵称": f"1 {group}"}]}
        with patch("app.core.order_merge_workflow.read_merge_orders", side_effect=orders), patch(
            "app.core.order_merge_workflow.get_wechat_group_members", side_effect=members
        ) as fetch:
            self.assertIn("总计2人", self.send("输出合发表"))
            self.assertEqual(fetch.call_count, 2)
            self.assertIn("总计2人", self.send("输出合发表"))
            self.assertEqual(fetch.call_count, 2)
        ctx = self.service.tools.get_context(self.session_id)
        workbook = load_workbook(ctx.merge_output_files[0])
        try:
            self.assertEqual(len(workbook.sheetnames), 3)
            self.assertEqual(workbook.worksheets[0].max_row, 3)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
