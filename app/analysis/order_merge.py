"""从原始订单读取身份和收货信息，生成互不重复的跨车合发清单。"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import os
import re
from tempfile import TemporaryDirectory, NamedTemporaryFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.analysis import order_parser as parser
from app.analysis.order_identity import resolve_order_identities
from app.utils.csv_utils import read_csv_dict_rows


SHIPPING_HEADERS = ("收货人姓名", "收货人联系方式", "收货人地址")


@dataclass
class MergeResult:
    path: str
    total: int
    combined: int
    counts: dict[str, int]
    warnings: list[str]


def read_merge_orders(path: str | Path, group: str) -> list[dict]:
    """每次从原文件解析商品，不依赖会话中可能过期的 parsed_orders。"""
    with TemporaryDirectory() as directory:
        parsed_path = Path(directory) / "parsed_orders.csv"
        parser.parse_order_file(path, parsed_path)
        products_by_serial = {}
        for row in read_csv_dict_rows(parsed_path)[0]:
            serial = str(row["单号"])
            if serial in products_by_serial:
                raise parser.OrderParseError(f"{group}存在重复单号：{serial}")
            products_by_serial[serial] = {
                key: int(value) for key, value in row.items()
                if key not in {"单号", "昵称", "总金额"} and value
            }

    wb = load_workbook(path, data_only=True)
    try:
        ws = wb.worksheets[parser.SHEET_INDEX]
        merged = parser._build_merged_value_map(ws)
        columns = {
            header: parser._find_required_header_col(ws, header, merged)
            for header in SHIPPING_HEADERS
        }
        orders = []
        for index in range(parser.HEADER_ROW + 1, ws.max_row + 1):
            if parser._is_empty_data_row(ws, index, ws.max_column, merged):
                continue
            serial = str(parser._parse_positive_int_required(
                parser._get_cell_value(ws, index, parser.ORDER_NO_COL, merged), "单号", index))
            raw_name = parser._get_cell_value(ws, index, parser.NICKNAME_COL, merged)
            name = "" if raw_name is None else str(raw_name)
            if not name.strip():
                raise parser.OrderParseError(f"{group}第{index}行订单昵称为空，已中止合发计算")
            shipping = tuple(
                str(value).strip() if value is not None else ""
                for value in (parser._get_cell_value(ws, index, columns[h], merged) for h in SHIPPING_HEADERS)
            )
            order = dict(name=name, serial=serial, row=index, shipping=shipping,
                         products=products_by_serial[serial])
            orders.append(order)
        return orders
    finally:
        wb.close()


def _content(order: dict | None) -> str:
    return "，".join(f"{name}×{quantity}" for name, quantity in order["products"].items()) if order else ""


def _sheet_name(group: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]\x00-\x1f]", "_", group).strip("'") or "车群"
    suffix = "补邮清单"
    base = cleaned[:31 - len(suffix)] + suffix
    name, number = base, 2
    while name.casefold() in used:
        tail = f"_{number}"
        name = base[:31 - len(tail)] + tail
        number += 1
    used.add(name.casefold())
    return name


def _append(ws, values):
    ws.append(values)
    # 原始昵称、商品及地址可能以等号开头，始终作为文本输出。
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def merge_order_files(sources: list[tuple[str, str | Path]], output: str | Path,
                      members_by_group: dict[str, list[dict]], *, resolved_orders_by_group=None) -> MergeResult:
    groups = [group for group, _ in sources]
    if len(groups) < 2 or len(set(groups)) != len(groups):
        raise ValueError("请提供至少两个不同车名")
    members: dict[str, dict[str, dict]] = {}
    identities = {}
    issues_by_group = {group: [] for group in groups}
    same_car_issues = {}
    counts = dict.fromkeys(groups, 0)
    for group, path in sources:
        try:
            if resolved_orders_by_group is not None:
                orders = resolved_orders_by_group[group]
            else:
                if group not in members_by_group:
                    raise ValueError("缺少微信成员身份数据")
                orders = resolve_order_identities(read_merge_orders(path, group), members_by_group[group], group)
        except (OSError, ValueError, RuntimeError) as error:
            raise parser.OrderParseError(f"车群“{group}”原始订单读取失败：{error}") from error
        for order in orders:
            wxid = order["wxid"]
            previous = identities.get(wxid)
            if previous and previous[1]["name"] != order["name"]:
                previous_group, previous_order = previous
                raise parser.OrderParseError(
                    f"同一 wxid（{wxid}）对应不同订单昵称："
                    f"{previous_group}单号{previous_order['serial']}“{previous_order['name']}”；"
                    f"{group}单号{order['serial']}“{order['name']}”")
            identities.setdefault(wxid, (group, order))
            group_orders = members.setdefault(wxid, {})
            if group in group_orders:
                aggregated = group_orders[group]
                aggregated["serial"] += "，" + order["serial"]
                if aggregated["shipping"] != order["shipping"]:
                    same_car_issues.setdefault(wxid, []).append(
                        f"{group}多单收货信息不一致（单号{aggregated['serial']}），"
                        f"该车以单号{aggregated['shipping_serial']}的收货信息为准")
                for product, quantity in order["products"].items():
                    aggregated["products"][product] = aggregated["products"].get(product, 0) + quantity
            else:
                group_orders[group] = dict(order, shipping_serial=order["serial"], products=dict(order["products"]))

    wb = Workbook()
    total_sheet = wb.active
    total_sheet.title = "总清单"
    sheets = {}
    used = {"总清单"}
    for group in groups:
        sheets[group] = wb.create_sheet(_sheet_name(group, used))

    def headers(order_groups, total=False):
        return (["wxid"] if total else []) + ["微信昵称", "订单昵称", "合发"] + [label for group in order_groups for label in
                (f"{group}单号", f"{group}订单内容")] + (["全部订单商品"] if total else []) + list(SHIPPING_HEADERS)

    _append(total_sheet, headers(groups, True))
    for group, ws in sheets.items():
        _append(ws, headers([group] + [other for other in groups if other != group]))

    combined = 0
    for wxid, orders in members.items():
        first = next(group for group in groups if group in orders)
        name = orders[first]["name"]
        wechat_name = orders[first]["wechat_name"]
        shipping_group = first
        if not orders[first]["shipping"][2]:
            shipping_group = next((group for group in groups if group in orders and orders[group]["shipping"][2]), first)
        shipping = orders[shipping_group]["shipping"]
        person = f"{wechat_name or '微信昵称未提供'}（{name}）"
        source = (f"本群（单号{orders[first]['shipping_serial']}）" if shipping_group == first else
                  f"车群{shipping_group}（单号{orders[shipping_group]['shipping_serial']}）")
        flag = len(orders) > 1
        combined += flag
        conflicts = [group for group, order in orders.items() if order["shipping"] != shipping]
        if shipping_group != first:
            issues_by_group[first].append(
                f"{person}本群（单号{orders[first]['shipping_serial']}）未填写收货地址，"
                f"收货信息不一致，暂以{source}的收货信息为准，请确认收货地址。")
        elif conflicts:
            issues_by_group[first].append(
                f"{person}收货信息不一致（涉及车群：{'、'.join([first] + conflicts)}），"
                f"以{source}的地址为准，请确认收货地址。")
        for issue in same_car_issues.get(wxid, []):
            issues_by_group[first].append(f"{person}{issue}；本清单以{source}的收货信息为准，请确认收货地址。")
        if any(not value for value in shipping):
            issues_by_group[first].append(f"{person}采用的{source}收货信息缺失：" +
                            "、".join(header for header, value in zip(SHIPPING_HEADERS, shipping) if not value) +
                            "，请补充并确认收货地址。")

        def values(order_groups):
            return [wechat_name, name, "是" if flag else "否"] + [value for group in order_groups for value in
                (orders[group]["serial"] if group in orders else "", _content(orders.get(group)))]

        all_products = "；".join(f"{group}-{_content(orders[group])}" for group in groups if group in orders)
        _append(total_sheet, [wxid] + values(groups) + [all_products] + list(shipping))
        _append(sheets[first], values([first] + [other for other in groups if other != first]) + list(shipping))
        counts[first] += 1

    for ws in wb:
        ws.freeze_panes = "E2" if ws is total_sheet else "D2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
        for column in ws.columns:
            header = str(column[0].value)
            ws.column_dimensions[column[0].column_letter].width = (
                48 if header in {"收货人地址", "全部订单商品"} else
                36 if header.endswith("订单内容") else 24 if header in {"订单昵称", "微信昵称", "wxid"} else 20)
            for cell in column:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    # 完整生成后再替换目标文件；错误不会留下半份工作簿。
    temporary = None
    try:
        with NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        wb.save(temporary)
        os.replace(temporary, output)
    finally:
        wb.close()
        if temporary is not None and temporary.exists():
            temporary.unlink()
    warnings = [f"车群{group}合发清单存在以下问题，请核对：\n" +
                "\n".join(f"{index}. {issue}" for index, issue in enumerate(issues, 1))
                for group, issues in issues_by_group.items() if issues]
    return MergeResult(str(output.resolve()), len(members), combined, counts, warnings)
